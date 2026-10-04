#!/usr/bin/env python3
"""Portable implementation of netmhcpan_ii_match_comparison. See the repository README and data/README.md for inputs, outputs, and execution instructions."""

from project_paths import DATA_ROOT

import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Implementation detail; see the repository documentation.
THRESH_EL  = 10.0    # Rank_EL    <= 10
THRESH_BA  = 10.0    # Rank_BA    <= 10
THRESH_AFF = 5000.0  # Affinity_nM < 5000

ANTIGENS = ["p24", "pp65", "PtxS1"]

SHEET_EL    = "Rank_EL_lte10"
SHEET_BA    = "Rank_BA_lte10"
SHEET_AFF   = "Affinity_lt5000"
SHEET_UNION = "Union_OR"

BASE_INPUT  = f"{DATA_ROOT}/{{ag}}/Matches MHC II/Matches NetMHCIIpan_4.1"
BASE_OUTPUT = f"{DATA_ROOT}/{{ag}}/Matches MHC II comparison"
INPUT_NAME  = "RM_{ag}_NetMHCIIpan_4.1.xlsx"
OUTPUT_NAME = "NetMHCIIpan_4.1_match_comparison.xlsx"

# Formatting.
HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
SUBHDR_FILL  = PatternFill("solid", fgColor="2E75B6")
SECTION_FILL = PatternFill("solid", fgColor="D6E4F0")
TOTAL_FILL   = PatternFill("solid", fgColor="E2EFDA")
WHITE_FILL   = PatternFill("solid", fgColor="FFFFFF")
ALT_FILL     = PatternFill("solid", fgColor="F2F7FC")

SUBHDR_FONT  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
BODY_FONT    = Font(name="Calibri", size=10)
BOLD_FONT    = Font(bold=True, name="Calibri", size=10)
SECTION_FONT = Font(bold=True, color="1F4E79", name="Calibri", size=10)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)


def thin_border():
    s = Side(style="thin", color="BDD7EE")
    return Border(left=s, right=s, top=s, bottom=s)


def style_cell(cell, fill=None, font=None, alignment=None):
    if fill:      cell.fill      = fill
    if font:      cell.font      = font
    if alignment: cell.alignment = alignment
    cell.border = thin_border()


# Input loading.
def load_rm(ag):
    path = os.path.join(
        BASE_INPUT.replace("{ag}", ag),
        INPUT_NAME.replace("{ag}", ag)
    )
    print(f"  Инпут файл: {path}")
    if not os.path.exists(path):
        print(f"  [ОШИБКА] Файл не найден: {path}")
        return None

    sheets = {}
    for sheet_name in (SHEET_EL, SHEET_BA, SHEET_AFF, SHEET_UNION):
        df = pd.read_excel(path, sheet_name=sheet_name)
        sheets[sheet_name] = df
        print(f"    лист «{sheet_name}»: {len(df)} строк")
    return sheets


# Threshold handling.
def analyse(sheets):
    df_el   = sheets[SHEET_EL]
    df_ba   = sheets[SHEET_BA]
    df_aff  = sheets[SHEET_AFF]
    df_uni  = sheets[SHEET_UNION]

    pairs = lambda df: set(zip(df["Allele"], df["Peptide"]))

    # Validation.
    for sheet_name, df in [
        (SHEET_EL,    df_el),
        (SHEET_BA,    df_ba),
        (SHEET_AFF,   df_aff),
        (SHEET_UNION, df_uni),
    ]:
        n_rows  = len(df)
        n_pairs = len(pairs(df))
        if n_rows != n_pairs:
            raise ValueError(
                f"В листе «{sheet_name}» найдены дубли пар Allele-Peptide: "
                f"{n_rows} строк, {n_pairs} уникальных пар."
            )

    p_el   = pairs(df_el)
    p_ba   = pairs(df_ba)
    p_aff  = pairs(df_aff)
    p_uni  = pairs(df_uni)

    n_total = len(p_uni)

    if n_total == 0:
        raise ValueError(
            "Лист Union_OR не содержит уникальных пар Allele-Peptide."
        )

    # Validation.
    computed = p_el | p_ba | p_aff
    if computed != p_uni:
        only_computed = computed - p_uni
        only_union    = p_uni - computed
        raise ValueError(
            f"Union_OR не совпадает с EL ∪ BA ∪ Aff. "
            f"Отсутствуют в Union_OR: {len(only_computed)}; "
            f"лишние в Union_OR: {len(only_union)}."
        )

    # Validation.
    # Implementation detail; see the repository documentation.
    # Validation.
    # Validation.
    # Implementation detail; see the repository documentation.
    flag_check = df_uni[["Allele", "Peptide"]].copy()
    flag_check["exp_el"]  = flag_check.apply(
        lambda r: (r["Allele"], r["Peptide"]) in p_el,  axis=1)
    flag_check["exp_ba"]  = flag_check.apply(
        lambda r: (r["Allele"], r["Peptide"]) in p_ba,  axis=1)
    flag_check["exp_aff"] = flag_check.apply(
        lambda r: (r["Allele"], r["Peptide"]) in p_aff, axis=1)
    flag_check["exp_cnt"] = (flag_check["exp_el"].astype(int) +
                             flag_check["exp_ba"].astype(int) +
                             flag_check["exp_aff"].astype(int))

    act_el  = df_uni["Pass_Rank_EL"].astype(bool)
    act_ba  = df_uni["Pass_Rank_BA"].astype(bool)
    act_aff = df_uni["Pass_Affinity"].astype(bool)
    act_cnt = df_uni["Passed_Thresholds_Count"].astype(int)

    bad_flags = (
        (act_el  != flag_check["exp_el"])  |
        (act_ba  != flag_check["exp_ba"])  |
        (act_aff != flag_check["exp_aff"]) |
        (act_cnt != flag_check["exp_cnt"])
    )
    if bad_flags.any():
        n_bad = bad_flags.sum()
        # Implementation detail; see the repository documentation.
        first_idx = bad_flags.idxmax()
        row_info  = df_uni.loc[first_idx]
        exp       = flag_check.loc[first_idx]
        raise ValueError(
            f"В Union_OR некорректные флаги у {n_bad} пар(ы). "
            f"Первая: ({row_info['Allele']}, {row_info['Peptide']}): "
            f"ожидалось EL={exp['exp_el']}, BA={exp['exp_ba']}, "
            f"Aff={exp['exp_aff']}, Count={exp['exp_cnt']}; "
            f"получено EL={bool(row_info['Pass_Rank_EL'])}, "
            f"BA={bool(row_info['Pass_Rank_BA'])}, "
            f"Aff={bool(row_info['Pass_Affinity'])}, "
            f"Count={int(row_info['Passed_Thresholds_Count'])}."
        )

    pct = lambda k: round(k / n_total * 100, 1) if n_total else 0.0

    # Implementation detail; see the repository documentation.
    el  = act_el
    ba  = act_ba
    aff = act_aff
    cnt = act_cnt

    # Implementation detail; see the repository documentation.
    n_single   = (cnt == 1).sum()
    n_two_plus = (cnt >= 2).sum()
    n_any      = n_total

    # Implementation detail; see the repository documentation.
    only_el  = ( el & ~ba & ~aff).sum()
    only_ba  = (~el &  ba & ~aff).sum()
    only_aff = (~el & ~ba &  aff).sum()
    el_ba    = ( el &  ba & ~aff).sum()
    el_aff   = ( el & ~ba &  aff).sum()
    ba_aff   = (~el &  ba &  aff).sum()
    all_three= ( el &  ba &  aff).sum()

    # Implementation detail; see the repository documentation.
    n_pass_el  = el.sum()
    n_pass_ba  = ba.sum()
    n_pass_aff = aff.sum()

    return {
        "n_total":    n_total,      "pct_total":    100.0,
        # Implementation detail; see the repository documentation.
        "n_single":   n_single,     "pct_single":   pct(n_single),
        "n_two_plus": n_two_plus,   "pct_two_plus": pct(n_two_plus),
        "n_any":      n_any,        "pct_any":      pct(n_any),
        # Implementation detail; see the repository documentation.
        "n_pass_el":  n_pass_el,    "pct_pass_el":  pct(n_pass_el),
        "n_pass_ba":  n_pass_ba,    "pct_pass_ba":  pct(n_pass_ba),
        "n_pass_aff": n_pass_aff,   "pct_pass_aff": pct(n_pass_aff),
        # Implementation detail; see the repository documentation.
        "n_only_el":  only_el,      "pct_only_el":  pct(only_el),
        "n_only_ba":  only_ba,      "pct_only_ba":  pct(only_ba),
        "n_only_aff": only_aff,     "pct_only_aff": pct(only_aff),
        "n_el_ba":    el_ba,        "pct_el_ba":    pct(el_ba),
        "n_el_aff":   el_aff,       "pct_el_aff":   pct(el_aff),
        "n_ba_aff":   ba_aff,       "pct_ba_aff":   pct(ba_aff),
        "n_all_three":all_three,    "pct_all_three":pct(all_three),
    }


# Output generation.
def write_excel(results, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение порогов"

    ags = [ag for ag in ANTIGENS if ag in results]
    n_ags = len(ags)

    ws.column_dimensions["A"].width = 50
    for i in range(n_ags * 2):
        ws.column_dimensions[get_column_letter(2 + i)].width = 10

    last_col = get_column_letter(1 + n_ags * 2)

    row = 1

    # Implementation detail; see the repository documentation.
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws.cell(row, 1, "Сравнение мэтчей NetMHCIIpan 4.1 по порогам")
    style_cell(c, fill=HEADER_FILL,
               font=Font(bold=True, color="FFFFFF", name="Calibri", size=13),
               alignment=CENTER)
    ws.row_dimensions[row].height = 24
    row += 1

    # Threshold handling.
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws.cell(row, 1,
        f"Пороги: EL %Rank <= {THRESH_EL}  |  BA %Rank <= {THRESH_BA}"
        f"  |  Affinity_nM < {THRESH_AFF:.0f}")
    style_cell(c, fill=SUBHDR_FILL,
               font=Font(italic=True, color="FFFFFF", name="Calibri", size=10),
               alignment=CENTER)
    row += 1

    # Implementation detail; see the repository documentation.
    hdr1, hdr2 = row, row + 1

    ws.merge_cells(f"A{hdr1}:A{hdr2}")
    c = ws.cell(hdr1, 1, "Показатель")
    style_cell(c, fill=SUBHDR_FILL, font=SUBHDR_FONT, alignment=CENTER)

    for col_i, ag in enumerate(ags):
        ag_col = 2 + col_i * 2
        ws.merge_cells(
            f"{get_column_letter(ag_col)}{hdr1}:{get_column_letter(ag_col+1)}{hdr1}"
        )
        c = ws.cell(hdr1, ag_col, ag)
        style_cell(c, fill=SUBHDR_FILL, font=SUBHDR_FONT, alignment=CENTER)

    for col_i in range(n_ags):
        ag_col = 2 + col_i * 2
        for lbl, ci in [("n", ag_col), ("%", ag_col + 1)]:
            c = ws.cell(hdr2, ci, lbl)
            style_cell(c, fill=SUBHDR_FILL, font=SUBHDR_FONT, alignment=CENTER)

    row = hdr2 + 1

    # Helper functions.
    def data_row(label, key_n, key_pct, section=False, total=False, alt=False):
        nonlocal row
        fill = SECTION_FILL if section else (TOTAL_FILL if total else
                                             (ALT_FILL if alt else WHITE_FILL))
        font = SECTION_FONT if section else (BOLD_FONT if total else BODY_FONT)
        c = ws.cell(row, 1, label)
        style_cell(c, fill=fill, font=font, alignment=LEFT)
        for col_i, ag in enumerate(ags):
            ag_col = 2 + col_i * 2
            r = results.get(ag, {})
            vn = r.get(key_n,  "—")
            vp = r.get(key_pct, "—")
            cn = ws.cell(row, ag_col,     vn)
            cp = ws.cell(row, ag_col + 1, f"{vp}%" if isinstance(vp, float) else vp)
            style_cell(cn, fill=fill, font=font, alignment=CENTER)
            style_cell(cp, fill=fill, font=font, alignment=CENTER)
        row += 1

    def section_header(label):
        nonlocal row
        ws.merge_cells(f"A{row}:{last_col}{row}")
        c = ws.cell(row, 1, label)
        style_cell(c, fill=SECTION_FILL, font=SECTION_FONT, alignment=LEFT)
        row += 1

    # Implementation detail; see the repository documentation.
    data_row("Всего уникальных мэтчей в Union_OR",
             "n_total", "pct_total", total=True)

    # Implementation detail; see the repository documentation.
    section_header("1. Мэтчи, проходящие ровно по 1 порогу")
    data_row("   Ровно 1 порог", "n_single", "pct_single", alt=True)

    # Implementation detail; see the repository documentation.
    section_header("2. Мэтчи, проходящие по >= 2 порогам")
    data_row("   >= 2 порогов", "n_two_plus", "pct_two_plus", alt=True)

    # Implementation detail; see the repository documentation.
    section_header("3. Все мэтчи, проходящие хотя бы по 1 порогу (Union_OR)")
    data_row("   >= 1 порога (итого)", "n_any", "pct_any", total=True)

    # Implementation detail; see the repository documentation.
    section_header("4. Распределение мэтчей по режимам предикции")

    data_row("   EL  (Rank_EL <= 10)  — всего",
             "n_pass_el", "pct_pass_el")
    data_row("   BA  (Rank_BA <= 10)  — всего",
             "n_pass_ba", "pct_pass_ba")
    data_row("   Aff (Affinity_nM < 5000)  — всего",
             "n_pass_aff", "pct_pass_aff")

    section_header("   Разбивка по непересекающимся группам (диаграмма Венна):")
    data_row("      Только EL  (не BA, не Aff)",
             "n_only_el",  "pct_only_el",  alt=True)
    data_row("      Только BA  (не EL, не Aff)",
             "n_only_ba",  "pct_only_ba",  alt=True)
    data_row("      Только Aff  (не EL, не BA)",
             "n_only_aff", "pct_only_aff", alt=True)
    data_row("      EL + BA  (не Aff)",
             "n_el_ba",    "pct_el_ba",    alt=True)
    data_row("      EL + Aff  (не BA)",
             "n_el_aff",   "pct_el_aff",   alt=True)
    data_row("      BA + Aff  (не EL)",
             "n_ba_aff",   "pct_ba_aff",   alt=True)
    data_row("      Все три (EL + BA + Aff)",
             "n_all_three","pct_all_three", total=True)

    # Implementation detail; see the repository documentation.
    row += 1
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws.cell(row, 1, "Примечания:")
    style_cell(c, fill=HEADER_FILL,
               font=Font(bold=True, color="FFFFFF", name="Calibri", size=10),
               alignment=LEFT)
    row += 1
    notes = [
        "% рассчитан от общего числа уникальных мэтчей в Union_OR для данного антигена.",
        "Порог EL: Rank_EL <= 10 (лист Rank_EL_lte10).",
        "Порог BA: Rank_BA <= 10 (лист Rank_BA_lte10).",
        "Порог Aff: Affinity_nM < 5000 нМ (лист Affinity_lt5000).",
        "Мэтч идентифицируется парой (Allele, Peptide).",
        "Булевы флаги Pass_Rank_EL/Pass_Rank_BA/Pass_Affinity берутся из листа Union_OR.",
        "Целостность проверяется: Union_OR должен совпадать с EL ∪ BA ∪ Aff.",
        "Группы раздела 4 непересекающиеся: каждый мэтч входит ровно в одну группу.",
    ]
    for note in notes:
        ws.merge_cells(f"A{row}:{last_col}{row}")
        c = ws.cell(row, 1, f"* {note}")
        style_cell(c, fill=ALT_FILL,
                   font=Font(italic=True, name="Calibri", size=9, color="1F4E79"),
                   alignment=LEFT)
        row += 1

    for r in range(4, row):
        ws.row_dimensions[r].height = 18

    wb.save(out_path)
    print(f"  Сохранено: {out_path}")


# ──────────────────────────── main ────────────────────────────────────────────
def main():
    print("=" * 60)
    print("NetMHCIIpan 4.1 — сравнение мэтчей по порогам")
    print("=" * 60)

    all_results = {}

    for ag in ANTIGENS:
        print(f"\n[{ag}] Загрузка данных...")
        sheets = load_rm(ag)
        if sheets is None:
            print(f"  Пропуск антигена {ag}: файл не найден.")
            continue

        try:
            result = analyse(sheets)
        except ValueError as e:
            print(f"  [ОШИБКА] {e}")
            print(f"  Пропуск антигена {ag}: данные не прошли проверку целостности.")
            continue

        all_results[ag] = result

        print(f"  Всего мэтчей (Union_OR): {result['n_total']}")
        print(f"  Ровно 1 порог: {result['n_single']} ({result['pct_single']}%)")
        print(f"  >= 2 порогов:  {result['n_two_plus']} ({result['pct_two_plus']}%)")
        print(f"  >= 1 порога:   {result['n_any']} ({result['pct_any']}%)")
        print(f"  EL (любой):    {result['n_pass_el']} ({result['pct_pass_el']}%)")
        print(f"  BA (любой):    {result['n_pass_ba']} ({result['pct_pass_ba']}%)")
        print(f"  Aff (любой):   {result['n_pass_aff']} ({result['pct_pass_aff']}%)")
        print(f"  Все три:       {result['n_all_three']} ({result['pct_all_three']}%)")

        out_dir = BASE_OUTPUT.replace("{ag}", ag)
        os.makedirs(out_dir, exist_ok=True)
        out_path = os.path.join(out_dir, OUTPUT_NAME)
        write_excel({ag: result}, out_path)

    if all_results:
        print(f"\n[Сводный файл] Запись таблицы для всех антигенов...")
        first_ag = next(ag for ag in ANTIGENS if ag in all_results)
        summary_dir = BASE_OUTPUT.replace("{ag}", first_ag)
        os.makedirs(summary_dir, exist_ok=True)
        summary_path = os.path.join(summary_dir,
                                    "NetMHCIIpan_4.1_match_comparison_ALL.xlsx")
        write_excel(all_results, summary_path)

    print(f"\n{'='*60}")
    print("Готово!")


if __name__ == "__main__":
    main()
